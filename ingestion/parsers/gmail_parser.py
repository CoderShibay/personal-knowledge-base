import mailbox
import os
import io
from datetime import datetime

try:
    import fitz
    PDF_SUPPORT = True
except ImportError:
    PDF_SUPPORT = False

try:
    from docx import Document
    DOCX_SUPPORT = True
except ImportError:
    DOCX_SUPPORT = False

try:
    import openpyxl
    XLSX_SUPPORT = True
except ImportError:
    XLSX_SUPPORT = False

# ── CONFIGURATION ──────────────────────────────────────────

ACCOUNTS = {
    "purno230":       "~/personal-kb/data/purno230",
    "ciai":           "~/personal-kb/data/ciai",
    "alisyed_office": "~/personal-kb/data/alisyed_office",
    "purnoli230":     "~/personal-kb/data/purnoli230",
    "purno240":       "~/personal-kb/data/purno240",
    "uni_aiub":       "~/personal-kb/data/uni_aiub",
}

SKIP_EXTENSIONS = [
    ".mp4", ".mkv", ".avi", ".mov",
    ".mp3", ".wav", ".flac",
    ".exe", ".dmg", ".iso", ".torrent"
]

IMAGE_EXTENSIONS = [".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".heic"]
ZIP_EXTENSIONS   = [".zip", ".rar", ".tar", ".gz"]

MAX_EMAILS_PER_ACCOUNT = None


# ── ATTACHMENT HANDLER ─────────────────────────────────────

def handle_attachment(part):
    """
    Returns a tuple: (text, attachment_tag)
    - text: any text we can extract right now (PDFs, docs)
    - attachment_tag: a note about files we'll process later
    """
    filename = part.get_filename() or ""
    ext = os.path.splitext(filename)[1].lower()

    # skip useless files silently
    if ext in SKIP_EXTENSIONS:
        return "", None

    try:
        payload = part.get_payload(decode=True)
        if not payload:
            return "", None
    except:
        return "", None

    # images — just tag for later, don't try to read pixels
    if ext in IMAGE_EXTENSIONS:
        return "", {"type": "image", "filename": filename}

    # zips — just tag for later, drive parser will handle
    if ext in ZIP_EXTENSIONS:
        return "", {"type": "zip", "filename": filename}

    # PDFs — extract text now
    if ext == ".pdf" and PDF_SUPPORT:
        try:
            doc = fitz.open(stream=payload, filetype="pdf")
            text = "".join(page.get_text() for page in doc)
            return f"\n[PDF: {filename}]\n{text}", None
        except:
            return "", {"type": "pdf_failed", "filename": filename}

    # Word docs — extract text now
    if ext == ".docx" and DOCX_SUPPORT:
        try:
            doc = Document(io.BytesIO(payload))
            text = "\n".join([p.text for p in doc.paragraphs])
            return f"\n[Word doc: {filename}]\n{text}", None
        except:
            return "", {"type": "docx_failed", "filename": filename}

    # Excel — extract text now
    if ext in [".xlsx", ".xls"] and XLSX_SUPPORT:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True)
            text = ""
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    row_text = " | ".join(str(c) for c in row if c is not None)
                    if row_text:
                        text += row_text + "\n"
            return f"\n[Excel: {filename}]\n{text}", None
        except:
            return "", {"type": "xlsx_failed", "filename": filename}

    # plain text
    if ext in [".txt", ".csv", ".md"]:
        try:
            return f"\n[Text file: {filename}]\n{payload.decode('utf-8', errors='ignore')}", None
        except:
            return "", None

    return "", None


# ── EMAIL BODY EXTRACTOR ───────────────────────────────────

def extract_email_body(message):
    body = ""
    attachment_texts = ""
    deferred_attachments = []  # images and zips saved for later

    if message.is_multipart():
        for part in message.walk():
            content_type = part.get_content_type()
            disposition  = str(part.get("Content-Disposition") or "")

            # main body text
            if content_type == "text/plain" and "attachment" not in disposition:
                try:
                    body += part.get_payload(decode=True).decode("utf-8", errors="ignore")
                except:
                    pass

            # attachments
            elif "attachment" in disposition or part.get_filename():
                text, tag = handle_attachment(part)
                attachment_texts += text
                if tag:
                    deferred_attachments.append(tag)
    else:
        try:
            body = message.get_payload(decode=True).decode("utf-8", errors="ignore")
        except:
            pass

    return body.strip(), attachment_texts.strip(), deferred_attachments


# ── MAIN PARSER ────────────────────────────────────────────

def parse_gmail_account(account_name, account_path):
    account_path = os.path.expanduser(account_path)
    chunks = []
    mbox_files = []

    for root, dirs, files in os.walk(account_path):
        for file in files:
            if file.endswith(".mbox"):
                mbox_files.append(os.path.join(root, file))

    if not mbox_files:
        print(f"  No .mbox files found in {account_path}")
        return chunks

    for mbox_path in mbox_files:
        print(f"  Reading: {os.path.basename(mbox_path)}")
        mbox = mailbox.mbox(mbox_path)
        count = 0

        for message in mbox:
            if MAX_EMAILS_PER_ACCOUNT and count >= MAX_EMAILS_PER_ACCOUNT:
                break

            subject = str(message.get("subject") or "No subject")
            sender  = str(message.get("from")    or "Unknown sender")
            date    = str(message.get("date")    or "Unknown date")
            to      = str(message.get("to")      or "")

            body, attachment_texts, deferred = extract_email_body(message)

            if not body and not attachment_texts:
                continue

            full_text = f"""Email from: {sender}
To: {to}
Date: {date}
Subject: {subject}

{body}
{attachment_texts}""".strip()

            # count how many images/zips are deferred
            image_files = [d["filename"] for d in deferred if d["type"] == "image"]
            zip_files   = [d["filename"] for d in deferred if d["type"] == "zip"]

            chunk = {
                "text": full_text,
                "metadata": {
                    "source":          "gmail",
                    "account":         account_name,
                    "subject":         subject,
                    "sender":          sender,
                    "date":            date,
                    "has_images":      bool(image_files),
                    "image_files":     image_files,
                    "has_zips":        bool(zip_files),
                    "zip_files":       zip_files,
                    "priority":        "normal"
                }
            }

            chunks.append(chunk)
            count += 1

        print(f"    Parsed {count} emails")

    return chunks


def parse_all_gmail_accounts():
    all_chunks = []

    for account_name, account_path in ACCOUNTS.items():
        print(f"\nProcessing account: {account_name}")
        chunks = parse_gmail_account(account_name, account_path)
        all_chunks.extend(chunks)
        print(f"  Total from {account_name}: {len(chunks)}")

    print(f"\nDone! Total emails parsed: {len(all_chunks)}")
    return all_chunks


if __name__ == "__main__":
    chunks = parse_all_gmail_accounts()

    if chunks:
        print("\n── PREVIEW OF FIRST CHUNK ──")
        print(f"Account:    {chunks[0]['metadata']['account']}")
        print(f"Subject:    {chunks[0]['metadata']['subject']}")
        print(f"Has images: {chunks[0]['metadata']['has_images']}")
        print(f"Has zips:   {chunks[0]['metadata']['has_zips']}")
        print(f"Preview:\n{chunks[0]['text'][:300]}...")
